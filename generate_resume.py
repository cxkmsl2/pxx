import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

doc = Document()

# 设置默认中文字体
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

head = doc.add_heading('简历项目经验：PXX 校园综合交易与服务平台', 0)
head.alignment = 1

doc.add_paragraph('技术栈：Go, Gin, gRPC, GORM, MySQL, Redis, Kafka, ETCD, Raft, Vue3, TypeScript')

doc.add_heading('项目描述：', level=2)
doc.add_paragraph('面向高校的高并发微服务交易系统，涵盖二手商品、闲置置换、拼团、租赁、数字拼卡及微跑腿等全方位校园服务。内置“全民小法庭”众裁系统解决交易纠纷，打造高内聚校园生态闭环。')

doc.add_heading('核心职责与技术亮点：', level=2)

p1 = doc.add_paragraph('1. 高可用微服务架构：')
p1.add_run('使用 Go-gRPC + Gin 构建 Gateway 及 4 大核心微服务（Account/Item/Trade/Feed）。基于 ETCD 实现服务注册与动态发现，保障内部调用高可用。')

p2 = doc.add_paragraph('2. TCC 分布式事务引擎：')
p2.add_run('针对“扣款+锁库存”等跨库交易，从零实现 TCC (Try-Confirm-Cancel) 模型。利用数据库唯一索引防悬挂，巧妙处理空回滚，并设计 Watchdog 定时拉起中间态事务，保障核心链路数据的绝对一致性。')

p3 = doc.add_paragraph('3. Raft + Kafka 延迟调度集群：')
p3.add_run('为解决海量“订单超时未支付自动取消”需求，使用 HashiCorp Raft 搭建多节点调度集群确保强一致性，通过 Kafka 削峰解耦，实现无单点故障的延迟任务分发状态机。')

p4 = doc.add_paragraph('4. 极速抢购与限流防刷：')
p4.add_run('拼团秒杀模块采用 Redis 分布式锁结合 MySQL 乐观锁 (`sold_count + quantity <= total_stock`)，彻底杜绝库存超卖；运用 Redis Pipeline 开发毫秒级滑动窗口限流中间件，抗击恶意洪峰。')

p5 = doc.add_paragraph('5. 并发状态机与 BFF 数据聚合：')
p5.add_run('在小法庭功能中利用悲观锁 (FOR UPDATE) 避免高并发投票的数据污染；在网关层完成跨库数据组装 (BFF 模式)，消除联表越界查询引发的级联崩溃。')

p6 = doc.add_paragraph('6. WebSocket 实时全双工通信：')
p6.add_run('基于 Gorilla WebSocket 开发高性能私信模块，通过读写协程分离与多端登录剔除机制治理连接泄漏；统一在网关做无状态 JWT 认证拦截，保障接口安全。')

doc.add_heading('性能压测表现：', level=2)
doc.add_paragraph('针对首页商品列表读取接口，在 WSL2 环境下使用 wrk 工具进行压测（参数：-t12 -c10000 -d10s），系统在 10,000 极端并发连接下，实现了单机约 9,500 QPS 的高吞吐量，且服务无宕机、无内存泄漏。')

doc.save(r'D:\go\Resume_PXX_V2.docx')
print('Docx updated successfully')
